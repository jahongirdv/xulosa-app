# -*- coding: utf-8 -*-
"""
Заменяет одиночные {xxx} на {{ xxx }} в shablon.docx.
Правит и параграфы, и таблицы, и колонтитулы.
"""
import re
from docx import Document


# Какие переменные должны быть
FIELDS = [
    "city", "inspection_date", "inspection_date_bold",
    "org_name", "org_str", "org_stir", "org_address", "org_head",
    "equipment_name", "equipment_model",
    "factory_number", "inventory_number", "year_manufactured",
    "lavozim", "fish", "bosh_hisobchi", "masul_shaxs",
]

# Регулярка ищет {xxx} где xxx — одно слово без пробелов и без {{ }}
PATTERN = re.compile(r'(?<!\{)\{([a-zA-Z_][a-zA-Z0-9_]*)\}(?!\})')


def fix_paragraph(p):
    """Заменяет {xxx} на {{ xxx }} в одном параграфе."""
    full = "".join(run.text for run in p.runs)
    new = PATTERN.sub(r'{{ \1 }}', full)

    if new != full:
        # сохраняем формат первого run
        if p.runs:
            first_run = p.runs[0]
            # оставляем один run с новым текстом
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            first_run.text = new
        else:
            p.add_run(new)


def fix_table(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                fix_paragraph(p)
            for nested in cell.tables:
                fix_table(nested)


def main(src, dst):
    doc = Document(src)

    # Параграфы
    for p in doc.paragraphs:
        fix_paragraph(p)

    # Таблицы
    for t in doc.tables:
        fix_table(t)

    # Колонтитулы
    for section in doc.sections:
        for p in section.header.paragraphs:
            fix_paragraph(p)
        for p in section.footer.paragraphs:
            fix_paragraph(p)

    doc.save(dst)
    print(f"✅ Готово: {dst}")


if __name__ == "__main__":
    import sys
    src = sys.argv[1] if len(sys.argv) > 1 else "shablon.docx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "shablon_fixed.docx"
    main(src, dst)
